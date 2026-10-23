"""Semantic facts, complete tokens, and server bypass/upgrade checks for the 1.4 report."""
import base64
import hashlib
import io
from pathlib import Path
from decimal import Decimal
import fitz
import pytest
from openpyxl import load_workbook
from backend.requirements import understand
from backend.requirement_tokens import quantities
from backend.catalog import UNITS
from backend.store import Store

TEXT='Need 20 bags cement for plastering. Prefer fast setting. Deliver to Shah Alam.'
CASES=[
 ('Need 1,020 bags cement 50kg. Deliver to Shah Alam.','quantity'),
 ('Need 20 bags cement for plastering and 20 bags nails. Deliver to Shah Alam.','two'),
 ('Need 20 bags cement for plastering, not fast setting. Deliver to Shah Alam.','not_fast'),
 ('我要20袋水泥用于抹灰，不要速凝，送到莎阿南。','not_fast'),
 ('Need 20 bags cement for plastering. No Brand A please. Deliver to Shah Alam.','not_brand'),
 ('Need 20 bags cement 50kg. Do not pick up; deliver to Shah Alam.','delivery'),
 ('我要20袋水泥50kg，不要自提，送到莎阿南。','delivery'),
 ('Need 20 bags cement for plastering.\nPrefer fast setting.\nDeliver to Shah Alam.','continuation')]

@pytest.mark.parametrize('text,kind',CASES)
def test_report_eight_inputs_preserve_business_facts(business_data,text,kind):
    d=understand(text,business_data['products'],business_data)
    selected=[next((p for p in business_data['products'] if p['sku']==i['sku']),None) for i in d['items']]
    if kind=='quantity':
        assert d['items'][0]['quantity']=='1020' and d['status']=='NEEDS_INFO'
        assert 'stock' in d['items'][0]['issues'];assert all(i['quantity']!='20' for i in d['items'])
    else:
        assert d['status']=='READY',d
        if kind=='two':assert len(selected)==2 and {p['category'] for p in selected}=={'cement','nails'} and [i['quantity'] for i in d['items']]==['20','20']
        elif kind=='not_fast':assert all('fast_setting' not in p['preferences'] for p in selected)
        elif kind=='not_brand':assert all(p['brand']!='Brand A' for p in selected)
        elif kind=='continuation':assert len(selected)==1 and selected[0]['sku']=='CEM-A-FAST-50'
    assert d['requirement']['fulfilment']=='delivery' and d['requirement']['area']=='Shah Alam'

@pytest.mark.parametrize('raw,expected',[('1,000','1000'),('12,345','12345'),('1,000,000','1000000'),('20','20'),('二十','20'),('twenty','20'),('dua puluh','20'),('two hundred and twenty','220')])
def test_complete_numeric_tokens_have_original_offsets(raw,expected):
    text='Need '+raw+' bags cement 50kg'
    q=quantities(text,UNITS);assert len(q)==1 and Decimal(q[0]['value'])==Decimal(expected)
    assert q[0]['raw']==raw and text[q[0]['span'][0]:q[0]['span'][1]].endswith('bags')

@pytest.mark.parametrize('raw',['1,02','1,020,','1,020.5','1 020','1，020','NaN','Infinity'])
def test_malformed_numbers_never_fall_back_to_numeric_suffix(raw):
    q=quantities('Need '+raw+' bags cement',UNITS)
    assert len(q)==1 and q[0]['error']=='quantity_format' and q[0]['value'] is None

@pytest.mark.parametrize('raw',['-1','0','1.5','1000001'])
def test_counted_quantities_still_reject_invalid_values(raw):
    assert quantities('Need '+raw+' bags cement',UNITS)[0]['error']=='invalid_quantity'

@pytest.mark.parametrize('text,method',[
 ('Need 20 bags cement 50kg. Do not pick up; deliver to Shah Alam.','delivery'),
 ('我要20袋水泥50kg，不要自提，送到莎阿南。','delivery'),
 ('Mahu 20 beg simen 50kg. Jangan ambil sendiri; hantar ke Shah Alam.','delivery'),
 ('Need 20 bags cement 50kg. Do not deliver; self collect.','pickup'),
 ('我要20袋水泥50kg，不要配送，我自提。','pickup'),
 ('Mahu 20 beg simen 50kg. Jangan hantar; ambil sendiri.','pickup')])
def test_three_language_delivery_negation_pairs(business_data,text,method):
    d=understand(text,business_data['products'],business_data);assert d['status']=='READY',d
    assert d['requirement']['fulfilment']==method
    assert d['requirement']['area']==('Shah Alam' if method=='delivery' else '')

@pytest.mark.parametrize('text',[
 'Need 20 bags cement for plastering, not only fast setting. Deliver to Shah Alam.',
 'Need 20 bags cement for plastering, not not fast setting. Deliver to Shah Alam.',
 '我要20袋水泥用于抹灰，不是不要速凝，送到莎阿南。',
 'Need 20 bags cement 50kg, self collect and deliver to Shah Alam.'
])
def test_uncertain_or_conflicting_constraints_require_explicit_resolution(business_data,text):
    d=understand(text,business_data['products'],business_data)
    assert not d['can_quote'] and any(p['code'] in {'negation_ambiguous','conflicting_constraints'} for p in d['issue_details'])

def test_ambiguous_modifier_and_repeated_quantity_are_not_discarded(business_data):
    d=understand('Need 20 bags cement\nNeed 1 can paint\nWaterproof\nDeliver to Shah Alam.',business_data['products'],business_data)
    assert 'ambiguous_modifier' in d['requirement']['issues'] and not d['can_quote']
    repeated=understand('CEM-A-50 70 bags\nCEM-A-50 70 bags\nDeliver to Shah Alam.',business_data['products'],business_data)
    assert len(repeated['items'])==2 and 'stock' in repeated['requirement']['issues']
    assert len({q['span'][0] for i in repeated['items'] for q in i['quantity_mentions']})==2


def request_review(env,text=TEXT,items=None,**kwargs):
    c,boss,staff,app=env
    body={'inquiry':text,'items':items or [{'sku':'CEM-A-FAST-50','quantity':'20'}],'area':'Shah Alam','fulfilment':'delivery',**kwargs}
    r=c.post('/api/requirements/review',headers=staff,json=body)
    assert r.status_code==200,r.text
    return r.json(),body


def draft(review,body,key='review-create-001'):
    return {'customer':'Synthetic review customer','inquiry':body['inquiry'],'items':[{k:i[k] for k in ('sku','quantity')} for i in body['items']],
            'language':'en','source_confirmed':True,'pricing_enabled':True,'delivery_area':body['area'],'fulfilment':body['fulfilment'],'request_key':key,'review_id':review['review_id']}


def test_review_required_for_direct_pricing_and_quote_creation(reviewed_app):
    c,_,staff,_=reviewed_app
    r=c.post('/api/pricing/preview',headers=staff,json={'inquiry':TEXT,'items':[{'sku':'CEM-A-FAST-50','quantity':'20'}],'area':'Shah Alam','fulfilment':'delivery'})
    assert r.status_code==422 and 'recommendation' not in r.json()
    r=c.post('/api/quotes',headers=staff,json={**draft({'review_id':None},{'inquiry':TEXT,'items':[{'sku':'CEM-A-FAST-50','quantity':'20'}],'area':'Shah Alam','fulfilment':'delivery'})})
    assert r.status_code in {409,422};assert c.get('/api/quotes',headers=staff).json()==[]

@pytest.mark.parametrize('change',[
 {'items':[{'sku':'CEM-A-FAST-50','quantity':'21'}]},
 {'items':[{'sku':'CEM-A-50','quantity':'20'}]},
 {'delivery_area':'Klang'}, {'fulfilment':'pickup','delivery_area':''}, {'inquiry':TEXT+' No Brand A please.'}
])
def test_review_id_cannot_authorize_changed_inputs(reviewed_app,change):
    c,_,staff,_=reviewed_app;r,b=request_review(reviewed_app)
    response=c.post('/api/quotes',headers=staff,json={**draft(r,b),**change})
    assert response.status_code==409,response.text
    assert c.get('/api/quotes',headers=staff).json()==[]


def test_direct_drop_and_negated_delivery_cannot_receive_review(reviewed_app):
    r,b=request_review(reviewed_app,CASES[1][0],[{'sku':'CEM-A-25','quantity':'20'}])
    assert not r['can_quote'] and any(i['code']=='missing_product' for i in r['issue_details'])
    r,_=request_review(reviewed_app,CASES[5][0],[{'sku':'CEM-A-50','quantity':'20'}],area='',fulfilment='pickup')
    assert not r['can_quote'] and any(i['code']=='fulfilment_changed' for i in r['issue_details'])


def test_concrete_quantity_correction_is_audited_and_exports_correct_value(reviewed_app):
    c,boss,staff,app=reviewed_app
    r,b=request_review(reviewed_app,items=[{'sku':'CEM-A-FAST-50','quantity':'21'}])
    p=next(p for p in r['issue_details'] if p['code']=='quantity_changed')
    correction={'issue_id':p['id'],'value':'21','reason':'Customer confirmed twenty-one bags'}
    r2,_=request_review(reviewed_app,**{**b,'corrections':[correction]})
    assert r2['can_quote'] and r2['resolutions'][0]['resolution']['value']=='21'
    payload=draft(r2,b);q=c.post('/api/quotes',headers=staff,json=payload).json()
    assert q['review']['resolutions'][0]['resolution']['reason']==correction['reason']
    assert c.post('/api/quotes',headers=staff,json=payload).json()['id']==q['id']
    assert c.post('/api/quotes/'+q['id']+'/submit',headers=staff,json={'expected_revision':1}).status_code==200
    done=c.post('/api/quotes/'+q['id']+'/approve',headers=boss,json={'expected_revision':1}).json()
    assert done['total']=='493.51' and '21' in done['approval']['replies']['en']
    pdf=c.get('/api/quotes/'+q['id']+'/pdf',headers=staff).content
    with fitz.open(stream=pdf,filetype='pdf') as doc:assert '493.51' in ''.join(p.get_text() for p in doc)
    assert any(e['action']=='REQUIREMENTS_REVIEWED' for e in done['activity'])


def test_stock_and_forged_correction_never_become_a_quote(reviewed_app):
    c,_,staff,_=reviewed_app
    r=c.post('/api/requirements/review',headers=staff,json={'inquiry':CASES[0][0],'items':[{'sku':'CEM-A-50','quantity':'1020'}],'area':'Shah Alam','fulfilment':'delivery'})
    assert r.status_code==422 and 'inventory' in r.text.lower()
    r,b=request_review(reviewed_app,items=[{'sku':'CEM-A-FAST-50','quantity':'21'}]);p=r['issue_details'][0]
    for correction in [{'issue_id':p['id'],'resolved':True,'reason':'verified'}, {'issue_id':p['id'],'value':'20','reason':'verified'}, {'issue_id':'invented','value':'21','reason':'verified'}]:
        bad=c.post('/api/requirements/review',headers=staff,json={**b,'corrections':[correction]});assert bad.status_code==422


def test_customer_cancelled_product_requires_specific_removal_record(reviewed_app):
    c,_,staff,_=reviewed_app;text=CASES[1][0]
    parsed=c.post('/api/inquiries/parse',headers=staff,json={'text':text,'mode':'requirements'}).json()
    items=[{'sku':parsed['items'][0]['sku'],'quantity':'20','source_id':parsed['items'][0]['source_id']}]
    r,b=request_review(reviewed_app,text,items);p=next(p for p in r['issue_details'] if p['code']=='missing_product')
    assert 'nails' in p['source_text']
    correction={'issue_id':p['id'],'value':{'action':'remove','source_id':p['source_id']},'reason':'Customer cancelled the nails'}
    accepted,_=request_review(reviewed_app,**{**b,'corrections':[correction]})
    assert accepted['can_quote'] and any(x['code']=='missing_product' for x in accepted['resolutions'])


def test_review_owner_context_versions_and_upgrade_snapshots(reviewed_app):
    c,boss,staff,app=reviewed_app;r,b=request_review(reviewed_app)
    assert c.post('/api/quotes',headers=boss,json=draft(r,b)).status_code==409
    q=c.post('/api/quotes',headers=staff,json=draft(r,b)).json()
    # Incremental migrations may run repeatedly without altering users or saved quotes.
    Store(app.state.store.directory);Store(app.state.store.directory)
    assert c.get('/api/quotes/'+q['id'],headers=staff).json()['id']==q['id']
    assert c.post('/api/quotes/'+q['id']+'/submit',headers=staff,json={'expected_revision':1}).status_code==200
    done=c.post('/api/quotes/'+q['id']+'/approve',headers=boss,json={'expected_revision':1}).json()
    before=c.get('/api/quotes/'+q['id']+'/pdf',headers=staff).content
    other,ob=request_review(reviewed_app);pending=c.post('/api/quotes',headers=staff,json=draft(other,ob,'pending-upgrade-002')).json()
    c.post('/api/quotes/'+pending['id']+'/submit',headers=staff,json={'expected_revision':1})
    with app.state.store.transaction() as db:db.execute("UPDATE requirement_reviews SET parser_version='requirements_rules_v1'")
    assert c.post('/api/quotes/'+pending['id']+'/approve',headers=boss,json={'expected_revision':1}).status_code==409
    assert c.get('/api/quotes/'+pending['id'],headers=staff).json()['review_required']
    assert c.get('/api/quotes/'+q['id']+'/pdf',headers=staff).content==before
    assert c.get('/api/quotes/'+q['id'],headers=staff).json()['total']==done['total']


def test_source_mapping_rejects_unknown_and_duplicate_segments(reviewed_app):
    c,_,staff,_=reviewed_app
    parsed=c.post('/api/inquiries/parse',headers=staff,json={'text':CASES[1][0],'mode':'requirements'}).json()
    rows=[{'sku':i['sku'],'quantity':i['quantity'],'source_id':i['source_id']} for i in parsed['items']]
    body={'inquiry':CASES[1][0],'items':rows,'area':'Shah Alam','fulfilment':'delivery'}
    for source in ['invented',rows[0]['source_id']]:
        changed=[rows[0],{**rows[1],'source_id':source}]
        assert c.post('/api/requirements/review',headers=staff,json={**body,'items':changed}).status_code==422


def test_review_is_scoped_to_revision_and_quote(reviewed_app):
    c,_,staff,_=reviewed_app
    a,ab=request_review(reviewed_app);q1=c.post('/api/quotes',headers=staff,json=draft(a,ab,'scope-quote-001')).json()
    b,bb=request_review(reviewed_app);q2=c.post('/api/quotes',headers=staff,json=draft(b,bb,'scope-quote-002')).json()
    changed=TEXT.replace('20 bags','22 bags')
    fresh,body=request_review(reviewed_app,changed,[{'sku':'CEM-A-FAST-50','quantity':'22'}],quote_id=q1['id'],expected_revision=1)
    revision={**draft(fresh,body),'expected_revision':1}
    assert c.put('/api/quotes/'+q2['id'],headers=staff,json=revision).status_code==409
    good=c.put('/api/quotes/'+q1['id'],headers=staff,json=revision)
    assert good.status_code==200 and good.json()['revision']==2
    assert c.post('/api/quotes/'+q1['id']+'/submit',headers=staff,json={'expected_revision':1}).status_code==409
    assert c.put('/api/quotes/'+q1['id'],headers=staff,json={**revision,'expected_revision':2}).status_code==409


def test_business_update_invalidates_unbound_review_and_prices(reviewed_app):
    c,boss,staff,_=reviewed_app;r,b=request_review(reviewed_app)
    book=load_workbook(Path(__file__).resolve().parents[1]/'samples/business_demo.xlsx')
    book['Rules'].cell(2,2,14);out=io.BytesIO();book.save(out)
    response=c.post('/api/business/import',headers=boss,json={'filename':'changed.xlsx','content_b64':base64.b64encode(out.getvalue()).decode()})
    assert response.status_code==200,response.text
    assert c.post('/api/quotes',headers=staff,json=draft(r,b)).status_code==409
    assert c.post('/api/pricing/preview',headers=staff,json={**b,'review_id':r['review_id']}).status_code==409


def test_pending_without_review_can_only_continue_after_new_reviewed_revision(reviewed_app):
    c,boss,staff,app=reviewed_app;r,b=request_review(reviewed_app)
    q=c.post('/api/quotes',headers=staff,json=draft(r,b)).json()
    c.post('/api/quotes/'+q['id']+'/submit',headers=staff,json={'expected_revision':1})
    with app.state.store.transaction() as db:db.execute('DELETE FROM quote_reviews WHERE quote_id=?',(q['id'],))
    assert c.get('/api/quotes/'+q['id'],headers=staff).json()['review_required']
    assert c.post('/api/quotes/'+q['id']+'/submit',headers=staff,json={'expected_revision':1}).status_code==409
    assert c.post('/api/quotes/'+q['id']+'/approve',headers=boss,json={'expected_revision':1}).status_code==409
    fresh,body=request_review(reviewed_app,quote_id=q['id'],expected_revision=1)
    revised=c.put('/api/quotes/'+q['id'],headers=staff,json={**draft(fresh,body),'expected_revision':1})
    assert revised.status_code==200 and revised.json()['revision']==2
    assert not revised.json()['review_required']
    assert c.post('/api/quotes/'+q['id']+'/submit',headers=staff,json={'expected_revision':2}).status_code==200
    assert c.post('/api/quotes/'+q['id']+'/approve',headers=boss,json={'expected_revision':2}).status_code==200
    with app.state.store.connect() as db:
        import json
        snap=json.loads(db.execute('SELECT snapshot FROM approvals WHERE quote_id=? AND revision=2',(q['id'],)).fetchone()[0])
        assert snap['requirement_review']['review_id']==fresh['review_id']


def test_uncertain_preference_requires_a_specific_choice_and_compatible_product(reviewed_app):
    c,_,staff,_=reviewed_app;text='Need 20 bags cement for plastering, not only fast setting. Deliver to Shah Alam.'
    r,b=request_review(reviewed_app,text)
    p=next(p for p in r['issue_details'] if p['code']=='negation_ambiguous')
    corrected={**b,'corrections':[{'issue_id':p['id'],'value':'include','reason':'Customer requires fast setting'}]}
    good=c.post('/api/requirements/review',headers=staff,json=corrected)
    assert good.status_code==200 and good.json()['can_quote']
    assert good.json()['effective']['items'][0]['constraints']['preferences']['include']==['fast_setting']
    incompatible=c.post('/api/requirements/review',headers=staff,json={**corrected,'items':[{'sku':'CEM-A-50','quantity':'20'}]})
    assert incompatible.status_code==200 and not incompatible.json()['can_quote']
    assert any(p['code']=='preference_changed' for p in incompatible.json()['issue_details'])


@pytest.mark.parametrize('text,kind',CASES[1:])
def test_each_supported_report_case_completes_review_approval_and_outputs(reviewed_app,text,kind):
    c,boss,staff,_=reviewed_app
    parsed=c.post('/api/inquiries/parse',headers=staff,json={'text':text,'mode':'requirements'}).json()
    assert parsed['can_quote']
    rows=[{k:i[k] for k in ('sku','quantity','source_id')} for i in parsed['items']]
    checked,body=request_review(reviewed_app,text,rows)
    assert checked['can_quote'],checked
    q=c.post('/api/quotes',headers=staff,json=draft(checked,body)).json()
    assert [{'sku':i['sku'],'quantity':i['quantity']} for i in q['items']]==[{'sku':i['sku'],'quantity':i['quantity']} for i in rows]
    assert c.post('/api/quotes/'+q['id']+'/submit',headers=staff,json={'expected_revision':1}).status_code==200
    response=c.post('/api/quotes/'+q['id']+'/approve',headers=boss,json={'expected_revision':1})
    assert response.status_code==200,response.text
    done=response.json();assert Decimal(done['pricing']['decision']['freight'])>0
    assert done['pricing']['requirement']['fulfilment']=='delivery' and done['pricing']['requirement']['area']=='Shah Alam'
    if kind=='two':assert len(done['items'])==2 and {i['sku'] for i in done['items']}=={'CEM-A-25','NAIL-A-1'}
    if kind=='not_fast':assert all(i['sku']!='CEM-A-FAST-50' for i in done['items'])
    if kind=='not_brand':assert all(i['brand']!='Brand A' for i in done['items'])
    if kind=='continuation':assert len(done['items'])==1 and done['items'][0]['sku']=='CEM-A-FAST-50'
    for reply in done['approval']['replies'].values():assert done['total'] in reply and done['pricing']['decision']['freight'] in reply
    with fitz.open(stream=c.get('/api/quotes/'+q['id']+'/pdf',headers=staff).content,filetype='pdf') as pdf:
        content=''.join(p.get_text() for p in pdf);assert done['total'] in content and done['pricing']['decision']['freight'] in content
    book=load_workbook(io.BytesIO(c.get('/api/exports/ledger',headers=staff).content))
    assert book['Items'].max_row==len(rows)+1 and book['Quotes'].cell(2,5).value==float(done['total'])
