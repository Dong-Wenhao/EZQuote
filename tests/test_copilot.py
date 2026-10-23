"""Full business workflow checks using synthetic inputs and a network-rejecting transport."""
import base64
import io
import json
from decimal import Decimal
from pathlib import Path

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app

ROOT = Path(__file__).resolve().parents[1]
TEXT = "Hi boss, I need 20 bags cement for plastering. Prefer fast setting. Deliver to Shah Alam. How much?"


class NoNetwork:
    def post(self, *args, **kwargs):
        raise AssertionError("The no-API workflow must not call an external provider")


def workbook_upload(change=None):
    raw = (ROOT / "samples/business_demo.xlsx").read_bytes()
    if change:
        book = load_workbook(io.BytesIO(raw))
        change(book)
        output = io.BytesIO(); book.save(output); raw = output.getvalue()
    return {"filename": "business_demo.xlsx", "content_b64": base64.b64encode(raw).decode()}


@pytest.fixture
def business(tmp_path):
    app = create_app(tmp_path, transport=NoNetwork(), start_worker=False)
    with TestClient(app) as c:
        setup = c.post("/api/setup", json={"shop_name": "Synthetic Supplies", "boss_password": "boss-test-123", "staff_password": "staff-test-123"}).json()
        boss = {"Authorization": "Bearer " + setup["token"]}
        staff = {"Authorization": "Bearer " + c.post("/api/login", json={"username": "staff", "password": "staff-test-123"}).json()["token"]}
        imported = c.post("/api/business/import", headers=boss, json=workbook_upload())
        assert imported.status_code == 200, imported.text
        yield c, boss, staff, app.state.store


def create(business, **overrides):
    c, boss, staff, store = business
    body = {"customer": "Synthetic Customer", "inquiry": TEXT, "language": "en", "source_confirmed": True,
            "items": [{"sku": "CEM-A-FAST-50", "quantity": "20"}], "request_key": "copilot-create-0001",
            "pricing_enabled": True, "delivery_area": "Shah Alam", **overrides}
    if "inquiry" not in overrides:
        if "items" in overrides:
            units = {p["sku"]: p["unit"] for p in c.get("/api/catalog", headers=staff).json()["products"]}
            body["inquiry"] = "; ".join(f"{i['sku']} {i['quantity']} {units[i['sku']]}" for i in body["items"]) + ". Deliver to Shah Alam."
        elif body.get("fulfilment") == "pickup":
            body["inquiry"] = "Need 20 bags fast setting cement for plastering, self collect."
    body["review_id"] = review_body(c, staff, body)["review_id"]
    r = c.post("/api/quotes", headers=staff, json=body)
    assert r.status_code == 200, r.text
    return r.json(), body


def review_body(c, auth, body, quote_id=None, revision=0):
    r = c.post("/api/requirements/review", headers=auth, json={
        "inquiry": body["inquiry"], "items": body["items"], "area": body.get("delivery_area", ""),
        "fulfilment": body.get("fulfilment", "delivery"), "quote_id": quote_id, "expected_revision": revision})
    assert r.status_code == 200, r.text
    assert r.json()["can_quote"], r.text
    return r.json()


def approve(business, q, **edits):
    c, boss, staff, store = business
    r = c.post(f"/api/quotes/{q['id']}/submit", headers=staff, json={"expected_revision": q["revision"]})
    assert r.status_code == 200, r.text
    r = c.post(f"/api/quotes/{q['id']}/approve", headers=boss, json={"expected_revision": q["revision"], **edits})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.mark.parametrize("text,language", [
    (TEXT, "en"),
    ("我要二十袋水泥用于抹灰，偏好速凝，送到莎阿南，多少钱？", "zh"),
    ("Saya mahu dua puluh beg simen untuk melepa, mahu cepat kering. Hantar ke Shah Alam. Berapa harga?", "ms")])
def test_brand_optional_requirement_matches_stock_and_use(business, text, language):
    c, boss, staff, _ = business
    r = c.post("/api/inquiries/parse", headers=staff, json={"text": text, "mode": "requirements"})
    assert r.status_code == 200
    data = r.json()
    assert data["status"] == "READY", data
    assert data["language"] == language
    assert data["method"] == "requirements_rules" and data["ai_state"] == "disabled"
    assert data["requirement"]["area"] == "Shah Alam"
    req = data["requirement"]["items"][0]
    assert req["brand"] is None and req["category"] == "cement"
    assert req["application"] == ["plastering"] and req["preference"] == ["fast_setting"]
    line = data["items"][0]
    assert line["sku"] == "CEM-A-FAST-50" and line["quantity"] == "20"
    assert line["best_match"]["stock"] == "85" and line["alternatives"]


@pytest.mark.parametrize("text,issue", [
    ("Need cement for plastering fast setting deliver to Shah Alam", "quantity"),
    ("Need -2 bags cement for plastering deliver to Shah Alam", "invalid_quantity"),
    ("Need 20 bags cement for plastering fast setting", "area"),
    ("Need 20 bags cement for plastering deliver to Unknown Town", "unknown_area"),
    ("Need Brand Z cement 50kg 20 bags deliver to Shah Alam", "unknown_product"),
    ("Need 20 bags cement 25kg or 50kg deliver to Shah Alam", "spec")])
def test_missing_or_conflicting_requirements_do_not_become_ready(business, text, issue):
    c, _, staff, _ = business
    data = c.post("/api/inquiries/parse", headers=staff, json={"text": text, "mode": "requirements"}).json()
    assert data["status"] == "NEEDS_INFO", data
    assert issue in data["requirement"]["issues"] + [x for item in data["items"] for x in item["issues"]]
    assert data["clarification"]


def test_pickup_does_not_require_area_or_add_freight(business):
    c, _, staff, _ = business
    data = c.post("/api/inquiries/parse", headers=staff, json={"text": "Need 20 bags fast setting cement for plastering, self collect", "mode": "requirements"}).json()
    assert data["status"] == "READY"
    assert data["requirement"]["fulfilment"] == "pickup" and data["requirement"]["area"] == ""
    q, _ = create(business, fulfilment="pickup", delivery_area="")
    r = q["pricing"]["recommendation"]
    assert r["base_freight"] == r["suggested_freight"] == r["estimated_delivery_cost"] == "0.00"


def test_preview_decimal_amounts_and_closed_denominator(business):
    c, _, staff, _ = business
    body = {"items": [{"sku": "CEM-A-FAST-50", "quantity": "20"}], "inquiry": TEXT, "delivery_area": "Shah Alam"}
    checked = review_body(c, staff, body)
    r = c.post("/api/pricing/preview", headers=staff, json={"items": body["items"], "inquiry": TEXT, "area": "Shah Alam", "fulfilment": "delivery", "review_id": checked["review_id"]})
    assert r.status_code == 200, r.text
    rec = r.json()["recommendation"]
    assert rec["original_product_price"] == "460.00" and rec["base_freight"] == "35.00" and rec["original_total"] == "495.00"
    assert rec["suggested_discount_pct"] == "3.00" and rec["suggested_discount_amount"] == "13.80"
    assert rec["suggested_freight"] == "25.00" and rec["recommended_total"] == "471.20"
    assert rec["similar_orders"] == 8 and rec["similar_won"] == 5 and rec["conversion_rate"] == "0.6250"
    assert rec["area_cost_orders"] == 5 and all(rec["reasons"][language] for language in ("zh", "en", "ms"))


def test_owner_edit_freezes_both_suggestion_and_final_outputs(business):
    c, boss, staff, store = business
    q, _ = create(business)
    # Staff cannot exercise the owner decision endpoint.
    assert c.post(f"/api/quotes/{q['id']}/approve", headers=staff, json={"expected_revision": 1}).status_code == 403
    done = approve(business, q, discount_pct="4", freight="30", pricing_note="Owner reviewed delivery")
    assert done["status"] == "COMPLETED" and done["total"] == "471.60"
    assert done["pricing"]["recommendation"]["recommended_total"] == "471.20"
    assert done["pricing"]["decision"]["mode"] == "edited"
    for reply in done["approval"]["replies"].values():
        assert "471.60" in reply and "30.00" in reply and "4.00" in reply
    pdf = c.get(f"/api/quotes/{q['id']}/pdf", headers=staff)
    document = fitz.open(stream=pdf.content, filetype="pdf")
    content = "".join(p.get_text() for p in document)
    assert "471.60" in content and "495.00" in content and "18.40" in content and "boss" in content
    assert "360.00" not in content  # Internal product cost is not printed on a customer quote.
    book = load_workbook(io.BytesIO(c.get("/api/exports/ledger", headers=staff).content))
    assert {"Pricing", "Requirements", "Quotes", "Items", "Activity"} <= set(book.sheetnames)
    row = list(book["Pricing"].iter_rows(min_row=2, values_only=True))[0]
    assert row[9] == 471.2 and row[13] == 471.6 and row[14] == "edited"
    before = len(done["activity"])
    repeat = c.post(f"/api/quotes/{q['id']}/approve", headers=boss, json={"expected_revision": 1, "discount_pct": "4", "freight": "30"})
    assert repeat.status_code == 200 and len(repeat.json()["activity"]) == before
    assert c.post(f"/api/quotes/{q['id']}/approve", headers=boss, json={"expected_revision": 1, "discount_pct": "3", "freight": "25"}).status_code == 409


@pytest.mark.parametrize("discount,freight", [("12", "30"), ("5", "0"), ("NaN", "25"), ("2", "-1")])
def test_discount_and_margin_policies_cannot_be_bypassed(business, discount, freight):
    c, boss, staff, _ = business
    q, _ = create(business)
    c.post(f"/api/quotes/{q['id']}/submit", headers=staff, json={"expected_revision": 1})
    r = c.post(f"/api/quotes/{q['id']}/approve", headers=boss, json={"expected_revision": 1, "discount_pct": discount, "freight": freight})
    assert r.status_code == 422, r.text
    unchanged = c.get(f"/api/quotes/{q['id']}", headers=staff).json()
    assert unchanged["approval"] is None and unchanged["pricing"]["decision"] is None


def test_inventory_is_checked_for_the_aggregate_sku_quantity(business):
    c, _, staff, _ = business
    body = {"customer": "Synthetic", "inquiry": TEXT, "language": "en", "source_confirmed": True,
            "items": [{"sku": "CEM-A-FAST-50", "quantity": "45"}, {"sku": "CEM-A-FAST-50", "quantity": "45"}],
            "request_key": "aggregate-stock-0001", "pricing_enabled": True, "delivery_area": "Shah Alam"}
    r = c.post("/api/quotes", headers=staff, json=body)
    assert r.status_code == 422 and "inventory" in r.text.lower()
    assert c.get("/api/quotes", headers=staff).json() == []


def test_client_cannot_disable_business_pricing_to_bypass_policy(business):
    c, _, staff, _ = business
    body = {"customer": "Synthetic", "inquiry": TEXT, "language": "en", "source_confirmed": True,
            "items": [{"sku": "CEM-A-FAST-50", "quantity": "20"}], "request_key": "no-policy-bypass-0001", "pricing_enabled": False}
    r = c.post("/api/quotes", headers=staff, json=body)
    assert r.status_code == 422 and "review is required" in r.text
    assert c.get("/api/quotes", headers=staff).json() == []


@pytest.mark.parametrize("change", [
    lambda b: b["Inventory"].append(["CEM-A-FAST-50", 85, 15, 80]),
    lambda b: b["History"].cell(2, 2, "2099-01-01"),
    lambda b: b["History"].cell(2, 6, "=1+2"),
    lambda b: b["Rules"].cell(2, 1, "unknown_rule"),
    lambda b: b["Inventory"].cell(2, 2, 2.5)])
def test_bad_business_workbooks_are_rejected_atomically(business, change):
    c, boss, _, _ = business
    old = c.get("/api/business", headers=boss).json()
    r = c.post("/api/business/import", headers=boss, json=workbook_upload(change))
    assert r.status_code == 422, r.text
    new = c.get("/api/business", headers=boss).json()
    assert new["id"] == old["id"] and new["catalog_id"] == old["catalog_id"]


def test_business_import_is_owner_only_and_idempotent(business):
    c, boss, staff, _ = business
    body = workbook_upload()
    for endpoint in ("/api/business/preview", "/api/business/import"):
        assert c.post(endpoint, headers=staff, json=body).status_code == 403
    old = c.get("/api/business", headers=boss).json()
    new = c.post("/api/business/import", headers=boss, json=body).json()
    assert old["id"] == new["id"] and old["catalog_id"] == new["catalog_id"]
    assert c.get("/copilot.js").status_code == 200


def test_business_changes_require_a_new_reviewed_revision(business):
    c, boss, staff, _ = business
    q, body = create(business)
    c.post(f"/api/quotes/{q['id']}/submit", headers=staff, json={"expected_revision": 1})
    changed = workbook_upload(lambda b: b["Inventory"].cell(8, 2, 18))
    assert c.post("/api/business/import", headers=boss, json=changed).status_code == 200
    assert c.post(f"/api/quotes/{q['id']}/approve", headers=boss, json={"expected_revision": 1}).status_code == 409
    current = c.get(f"/api/quotes/{q['id']}", headers=staff).json()
    assert current["pricing"]["recommendation"]["stock_levels"]["CEM-A-FAST-50"]["stock"] == "85"
    assert current["approval"] is None


def test_no_history_uses_conservative_defaults(business):
    c, boss, staff, _ = business
    body = workbook_upload(lambda b: b["History"].delete_rows(2, 100))
    assert c.post("/api/business/import", headers=boss, json=body).status_code == 200
    q, _ = create(business)
    rec = q["pricing"]["recommendation"]
    assert rec["similar_orders"] == 0 and rec["conversion_rate"] is None
    assert rec["suggested_discount_pct"] == "0.00" and rec["suggested_freight"] == "35.00" and rec["confidence"] == "low"


def test_recommendation_reduces_discount_to_keep_margin(business):
    c, boss, _, _ = business
    body = workbook_upload(lambda b: b["Products"].cell(8, 13, 19))
    assert c.post("/api/business/import", headers=boss, json=body).status_code == 200
    q, _ = create(business)
    rec = q["pricing"]["recommendation"]
    assert Decimal(rec["suggested_discount_pct"]) < 3
    assert Decimal(rec["estimated_margin_pct"]) >= 15


@pytest.mark.parametrize("actual_cost", [None, "29"])
def test_outcome_changes_future_statistics_and_costs_need_actual_evidence(business, actual_cost):
    c, boss, staff, _ = business
    q, _ = create(business)
    done = approve(business, q, discount_pct="4", freight="30")
    body = {"expected_revision": 1, "outcome": "won", "actual_delivery_cost": actual_cost}
    r = c.post(f"/api/quotes/{q['id']}/outcome", headers=staff, json=body)
    assert r.status_code == 200, r.text
    assert r.json()["outcome"]["version"] == 1
    repeat = c.post(f"/api/quotes/{q['id']}/outcome", headers=staff, json=body)
    assert repeat.status_code == 200 and repeat.json()["outcome"]["version"] == 1
    stats = c.get("/api/business", headers=boss).json()
    assert stats["live_orders"] == 1
    q2, _ = create(business, request_key="future-quote-0002")
    r2 = q2["pricing"]["recommendation"]
    assert r2["similar_orders"] == 9 and r2["similar_won"] == 6
    assert r2["historical_average_discount_pct"] == "3.17"
    assert r2["area_cost_orders"] == (6 if actual_cost else 5)
    # The first quotation retains the original recommendation after new outcomes.
    assert c.get(f"/api/quotes/{q['id']}", headers=staff).json()["pricing"]["recommendation"]["similar_orders"] == 8
    stale = c.post(f"/api/quotes/{q['id']}/outcome", headers=staff, json={"expected_revision": 1, "outcome": "lost"})
    assert stale.status_code == 409
    book = load_workbook(io.BytesIO(c.get("/api/exports/ledger", headers=boss).content))
    assert "Outcomes" in book.sheetnames and book["Outcomes"].cell(2, 3).value == "won"


def test_customer_outcomes_require_approval_and_revisions_are_one_order(business):
    c, boss, staff, _ = business
    q, body = create(business)
    assert c.post(f"/api/quotes/{q['id']}/outcome", headers=staff, json={"expected_revision": 1, "outcome": "won"}).status_code == 409
    approve(business, q)
    assert c.post(f"/api/quotes/{q['id']}/outcome", headers=staff, json={"expected_revision": 1, "outcome": "won"}).status_code == 200
    revised = {**body, "expected_revision": 1, "inquiry": TEXT.replace("20 bags", "22 bags"), "items": [{"sku": "CEM-A-FAST-50", "quantity": "22"}]}
    revised["review_id"] = review_body(c, staff, revised, q["id"], 1)["review_id"]
    r = c.put(f"/api/quotes/{q['id']}", headers=staff, json=revised)
    assert r.status_code == 200, r.text
    # A revised order must not train on its own previous revision.
    assert r.json()["pricing"]["recommendation"]["similar_orders"] == 8
    approve(business, r.json())
    assert c.post(f"/api/quotes/{q['id']}/outcome", headers=staff, json={"expected_revision": 2, "outcome": "won"}).status_code == 200
    assert c.get("/api/business", headers=boss).json()["live_orders"] == 1


def test_multi_item_orders_do_not_double_count_conversion_or_freight(business):
    c, boss, staff, _ = business
    def add_basket(book):
        h = book["History"]
        initial = list(h.iter_rows(min_row=2, values_only=True))
        for row in initial:
            if str(row[0]).startswith("DEMO-FAST"):
                h.append([*row[:2], "PVC-A-20", 5, *row[4:]])
    assert c.post("/api/business/import", headers=boss, json=workbook_upload(add_basket)).status_code == 200
    q, _ = create(business, items=[{"sku": "CEM-A-FAST-50", "quantity": "20"}, {"sku": "PVC-A-20", "quantity": "5"}])
    rec = q["pricing"]["recommendation"]
    assert rec["similar_orders"] == 8 and rec["similar_won"] == 5 and rec["area_cost_orders"] == 5
    assert rec["recommended_total"] == "513.64"


def test_upgrade_keeps_approved_legacy_files_but_requires_review_of_pending_quotes(tmp_path):
    app = create_app(tmp_path, transport=NoNetwork(), start_worker=False)
    with TestClient(app) as c:
        auth = {"Authorization": "Bearer " + c.post("/api/setup", json={"shop_name": "Old shop", "boss_password": "boss-upgrade-test", "staff_password": "staff-upgrade-test"}).json()["token"]}
        c.post("/api/catalog/import", headers=auth, json={"filename": "products.xlsx", "content_b64": base64.b64encode((ROOT / "samples/products.xlsx").read_bytes()).decode()})
        body = {"customer": "Legacy", "inquiry": "A牌50kg水泥20袋", "language": "zh", "source_confirmed": True, "items": [{"sku": "CEM-A-50", "quantity": "20"}], "request_key": "upgrade-old-0001"}
        old = c.post("/api/quotes", headers=auth, json=body).json()
        c.post(f"/api/quotes/{old['id']}/submit", headers=auth, json={"expected_revision": 1})
        assert c.post(f"/api/quotes/{old['id']}/approve", headers=auth, json={"expected_revision": 1}).status_code == 200
        pending = c.post("/api/quotes", headers=auth, json={**body, "request_key": "upgrade-pending-0002"}).json()
        c.post(f"/api/quotes/{pending['id']}/submit", headers=auth, json={"expected_revision": 1})
        c.post("/api/business/import", headers=auth, json=workbook_upload())
        assert c.get(f"/api/quotes/{old['id']}/pdf", headers=auth).status_code == 200
        assert c.get(f"/api/quotes/{old['id']}", headers=auth).json()["total"] == "370.00"
        assert c.get(f"/api/quotes/{pending['id']}", headers=auth).json()["review_required"]
        assert c.post(f"/api/quotes/{pending['id']}/approve", headers=auth, json={"expected_revision": 1}).status_code == 409
