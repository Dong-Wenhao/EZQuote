"""Complete expressions and unclassified coordinated goods from the 1.5.1 retest."""
import io
import fitz
import pytest
from openpyxl import load_workbook
from backend.catalog import UNITS,parse_inquiry
from backend.requirement_tokens import quantities
from backend.requirements import understand


REPORT_CASES=[
    ('Need 20／30 bags cement 50kg. Deliver to Shah Alam.','30','multiple_quantities'),
    ('Need 1_020 bags cement 50kg. Deliver to Shah Alam.','20','quantity_format'),
    ('Need 20 bags cement 50kg and sand. Deliver to Shah Alam.','20','missing_product'),
]
MALFORMED=['1_020','1__020','1..020','1,,020','1e3','20//30','20 ／／ 30',
           '20⁄30','20 ∕ 30','1\u200b020','1\u2009020','1²020','1点020','20%30','1e 020','20or 30','NaN and 20']
RANGES=['20／30','20 ／ 30','２０／３０','二十／三十','twenty／thirty','dua puluh／tiga puluh']


@pytest.mark.parametrize('raw',MALFORMED+RANGES)
def test_complete_expression_is_retained_with_original_offsets_and_no_numeric_suffix(business_data,raw):
    text='🙂 客户：Need '+raw+' bags cement 50kg'
    mentions=quantities(text,UNITS,offset=7)
    assert len(mentions)==1
    q=mentions[0]
    assert q['raw']==raw and q['value'] is None
    assert q['error'] in {'quantity_format','multiple_quantities'}
    assert text[q['span'][0]-7:q['span'][1]-7]==raw+' bags'
    parsed=parse_inquiry('CEM-A-50 '+raw+' bags',business_data['products'])
    assert parsed['status']=='NEEDS_INFO' and parsed['items'][0]['quantity']==''


@pytest.mark.parametrize('raw,value',[('twenty-one','21'),('dua-puluh','20'),('two hundred and twenty','220')])
def test_compound_word_quantities_still_work(raw,value):
    q=quantities('Need '+raw+' bags cement 50kg',UNITS)
    assert len(q)==1 and q[0]['raw']==raw and q[0]['value']==value and q[0]['error'] is None


@pytest.mark.parametrize('text,bad_quantity,expected_code',REPORT_CASES)
def test_report_inputs_cannot_receive_a_review_or_create_a_quote_even_with_old_wrong_values(reviewed_app,text,bad_quantity,expected_code):
    c,_,staff,_=reviewed_app
    parsed=c.post('/api/inquiries/parse',headers=staff,json={'text':text,'mode':'requirements'}).json()
    assert not parsed['can_quote'] and parsed['status']=='NEEDS_INFO'
    first=parsed['items'][0]
    rows=[{'sku':'CEM-A-50','quantity':bad_quantity,'source_id':first['source_id']}]
    body={'inquiry':text,'items':rows,'area':'Shah Alam','fulfilment':'delivery'}
    checked=c.post('/api/requirements/review',headers=staff,json=body)
    assert checked.status_code==200 and not checked.json()['can_quote']
    assert not checked.json().get('review_id')
    assert any(p['code']==expected_code for p in checked.json()['issue_details'])
    selected=[{k:i[k] for k in ('sku','quantity')} for i in rows]
    quote={'customer':'Synthetic boundary check','inquiry':text,'items':selected,'language':'en',
           'source_confirmed':True,'request_key':'boundary-check-001','pricing_enabled':True,
           'delivery_area':'Shah Alam','fulfilment':'delivery'}
    created=c.post('/api/quotes',headers=staff,json=quote)
    assert created.status_code==409 and created.json()['detail']['code']=='review_stale'
    preview=c.post('/api/pricing/preview',headers=staff,json={**body,'items':selected})
    assert preview.status_code==422 and preview.json()['detail']['code']=='review_required'
    assert c.get('/api/quotes',headers=staff).json()==[]


UNKNOWN_CASES=[
    ('Need 20 bags cement 50kg and sand. Deliver to Shah Alam.','sand'),
    ('Mahu 20 beg simen 50kg dan pasir. Hantar ke Shah Alam.','pasir'),
    ('我要20袋水泥50kg和沙子，送到莎阿南。','沙子'),
    ('Need 20 bags cement 50kg, sand. Deliver to Shah Alam.','sand'),
    ('Need sand and 20 bags cement 50kg. Deliver to Shah Alam.','sand'),
    ('Need 20 bags cement 50kg and waterproof sand. Deliver to Shah Alam.','sand'),
    ('Need 20 bags cement 50kg and deliver to Shah Alam and sand.','sand'),
    ('Need 20 bags cement 50kg.\nWaterproof sand. Deliver to Shah Alam.','sand'),
]


@pytest.mark.parametrize('text,unknown',UNKNOWN_CASES)
def test_unknown_coordinated_fragment_is_retained_even_with_logistics_or_modifiers(business_data,text,unknown):
    d=understand(text,business_data['products'],business_data)
    assert not d['can_quote'] and len(d['items'])==2
    known=next(i for i in d['items'] if i['sku'])
    assert known['sku']=='CEM-A-50' and known['quantity']=='20'
    unclear=next(i for i in d['items'] if unknown in i['source_text'])
    assert 'unknown_product' in unclear['issues'] and 'quantity' in unclear['issues']
    assert not unclear['sku'] and not unclear['quantity']
    for part in unclear['source_segments']:
        assert text[slice(*part['span'])]==part['source_text']


@pytest.mark.parametrize('continuation',[
    'and fast setting. Deliver to Shah Alam.',
    'and deliver to Shah Alam. How much?',
    'and please deliver to Shah Alam.',
])
def test_classified_preference_delivery_and_courtesy_are_not_phantom_goods(business_data,continuation):
    d=understand('Need 20 bags cement 50kg '+continuation,business_data['products'],business_data)
    assert d['can_quote'] and len(d['items'])==1
    assert d['items'][0]['quantity']=='20' and d['requirement']['area']=='Shah Alam'


def test_unknown_goods_cancellation_is_explicit_and_survives_approval_and_ledger(reviewed_app):
    c,boss,staff,_=reviewed_app;text=REPORT_CASES[2][0]
    parsed=c.post('/api/inquiries/parse',headers=staff,json={'text':text,'mode':'requirements'}).json()
    first=parsed['items'][0];rows=[{k:first[k] for k in ('sku','quantity','source_id')}]
    body={'inquiry':text,'items':rows,'area':'Shah Alam','fulfilment':'delivery'}
    pending=c.post('/api/requirements/review',headers=staff,json=body).json()
    p=next(p for p in pending['issue_details'] if p['code']=='missing_product')
    assert 'sand' in p['source_text']
    reason='Customer confirmed that the sand is cancelled'
    correction={'issue_id':p['id'],'value':p['expected_value'],'reason':reason}
    checked=c.post('/api/requirements/review',headers=staff,json={**body,'corrections':[correction]}).json()
    assert checked['can_quote']
    saved=c.post('/api/quotes',headers=staff,json={'customer':'Synthetic cancellation','inquiry':text,'items':[{k:i[k] for k in ('sku','quantity')} for i in rows],
        'language':'en','source_confirmed':True,'request_key':'cancel-unknown-001','pricing_enabled':True,
        'delivery_area':'Shah Alam','fulfilment':'delivery','review_id':checked['review_id']})
    assert saved.status_code==200,saved.text
    q=saved.json()
    assert c.post('/api/quotes/'+q['id']+'/submit',headers=staff,json={'expected_revision':1}).status_code==200
    approved=c.post('/api/quotes/'+q['id']+'/approve',headers=boss,json={'expected_revision':1})
    assert approved.status_code==200
    done=approved.json()
    assert done['status']=='COMPLETED' and done['total']=='398.00' and len(done['items'])==1
    assert any(r['code']=='missing_product' and r['resolution']['reason']==reason for r in done['review']['resolutions'])
    with fitz.open(stream=c.get('/api/quotes/'+q['id']+'/pdf',headers=staff).content,filetype='pdf') as doc:
        assert done['total'] in ''.join(page.get_text() for page in doc)
    book=load_workbook(io.BytesIO(c.get('/api/exports/ledger',headers=staff).content))
    assert book['Items'].max_row==2
    assert any(reason in str(cell.value) for row in book['Requirements'] for cell in row)


def test_unknown_fragment_needs_both_product_and_quantity_confirmation_when_replaced(reviewed_app):
    c,_,staff,_=reviewed_app;text=REPORT_CASES[2][0]
    parsed=c.post('/api/inquiries/parse',headers=staff,json={'text':text,'mode':'requirements'}).json()
    a,b=parsed['items']
    rows=[{k:a[k] for k in ('sku','quantity','source_id')},{'sku':'NAIL-A-1','quantity':'5','source_id':b['source_id']}]
    body={'inquiry':text,'items':rows,'area':'Shah Alam','fulfilment':'delivery'}
    pending=c.post('/api/requirements/review',headers=staff,json=body).json()
    problems=[p for p in pending['issue_details'] if p['source_id']==b['source_id']]
    assert {p['code'] for p in problems}=={'quantity','unknown_product'}
    corrections=[{'issue_id':p['id'],'value':p['expected_value'],'reason':'Customer replaces sand with five bags of nails'} for p in problems]
    partial=c.post('/api/requirements/review',headers=staff,json={**body,'corrections':corrections[:1]}).json()
    assert not partial['can_quote']
    checked=c.post('/api/requirements/review',headers=staff,json={**body,'corrections':corrections}).json()
    assert checked['can_quote'] and len(checked['resolutions'])==2
