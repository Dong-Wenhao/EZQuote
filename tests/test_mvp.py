import base64
import io
import json
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app
from backend.catalog import parse_inquiry

ROOT = Path(__file__).resolve().parents[1]
PRODUCTS = json.loads((ROOT / "samples/products.json").read_text("utf-8"))


@pytest.fixture
def env(tmp_path):
    app = create_app(tmp_path)
    with TestClient(app) as client:
        setup = client.post("/api/setup", json={"shop_name": "合成建材 Demo Supplies", "boss_password": "boss-test-123", "staff_password": "staff-test-123"})
        assert setup.status_code == 200, setup.text
        boss = {"Authorization": "Bearer " + setup.json()["token"]}
        login = client.post("/api/login", json={"username": "staff", "password": "staff-test-123"})
        staff = {"Authorization": "Bearer " + login.json()["token"]}
        upload = {"filename": "products.xlsx", "content_b64": base64.b64encode((ROOT / "samples/products.xlsx").read_bytes()).decode()}
        imported = client.post("/api/catalog/import", headers=boss, json=upload)
        assert imported.status_code == 200, imported.text
        yield client, boss, staff, app.state.store, upload


def draft(client, staff, key="test-key-001", quantity="20", customer="合成客户 01"):
    body = {"customer": customer, "inquiry": "A牌50kg水泥20袋，多少钱？", "language": "zh",
            "items": [{"sku": "CEM-A-50", "quantity": quantity}], "notes": "", "source_confirmed": True, "request_key": key}
    response = client.post("/api/quotes", headers=staff, json=body)
    assert response.status_code == 200, response.text
    return response.json(), body


def submit(client, staff, q):
    r = client.post("/api/quotes/" + q["id"] + "/submit", headers=staff, json={"expected_revision": q["revision"]})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.mark.parametrize("text,language", [
    ("A牌50kg水泥20袋，多少钱？", "zh"),
    ("How much for 20 bags of Brand A cement 50kg?", "en"),
    ("Berapa harga 20 beg simen Brand A 50kg?", "ms"),
    ("A牌 simen 50kg 要20 bag", "zh"),
    ("CEM-A-50 20 bags", "en"),
])
def test_three_languages(text, language):
    parsed = parse_inquiry(text, PRODUCTS)
    assert parsed["status"] == "READY", parsed
    assert parsed["language"] == language
    assert parsed["items"][0]["sku"] == "CEM-A-50"
    assert parsed["items"][0]["quantity"] == "20"


@pytest.mark.parametrize("text,issue", [
    ("A牌水泥20袋", "spec"), ("50kg水泥20袋", "brand"), ("A牌50kg水泥", "quantity"),
    ("A牌50kg水泥0袋", "invalid_quantity"), ("A牌50kg水泥-1袋", "invalid_quantity"),
    ("A牌50kg水泥1.5袋", "invalid_quantity"), ("Brand Z cement 50kg 20 bags", "unknown_product"),
    ("Brand A cement 50kg 20 bags and 10 bags", "multiple_quantities"),
])
def test_ambiguity_blocks(text, issue):
    parsed = parse_inquiry(text, PRODUCTS)
    assert parsed["status"] == "NEEDS_INFO"
    assert issue in parsed["items"][0]["issues"]
    assert parsed["items"][0]["sku"] == ""
    assert parsed["clarification"]


def test_multiline():
    p = parse_inquiry("A牌50kg水泥20袋\nBrand A PVC pipe 20mm 5 pieces", PRODUCTS)
    assert p["status"] == "READY", p
    assert [i["sku"] for i in p["items"]] == ["CEM-A-50", "PVC-A-20"]


def test_chat_source_is_per_revision_and_insert_log_is_not_send(env):
    c, boss, staff, store, _ = env
    _, body = draft(c, staff, key="source-base-001")
    source = {"kind": "whatsapp", "tab_id": 5, "page_session": "synthetic-page",
              "chat_token": "synthetic-chat", "chat_title": "Synthetic Customer",
              "message_id": "false_synthetic@c.us_MESSAGE", "anchor": "synthetic-anchor"}
    body.update(request_key="source-bound-001", source=source)
    q = c.post("/api/quotes", headers=staff, json=body).json()
    assert q["source"] == source
    assert c.get("/api/quotes/" + q["id"], headers=boss).json()["source"] == source
    assert c.post("/api/quotes/" + q["id"] + "/activity", headers=staff,
                  json={"action": "REPLY_INSERTED", "revision": 1}).status_code == 409
    q = submit(c, staff, q)
    assert c.post("/api/quotes/" + q["id"] + "/approve", headers=boss,
                  json={"expected_revision": 1}).status_code == 200
    assert c.post("/api/quotes/" + q["id"] + "/activity", headers=staff,
                  json={"action": "REPLY_INSERTED", "revision": 1}).status_code == 200
    revised = c.put("/api/quotes/" + q["id"], headers=staff,
                    json={**body, "source": None, "expected_revision": 1})
    assert revised.status_code == 200, revised.text
    assert revised.json()["source"] is None and revised.json()["approval"] is None
    with store.connect() as db:
        old_source = db.execute("SELECT source FROM quote_sources WHERE quote_id=? AND revision=1", (q["id"],)).fetchone()
        assert json.loads(old_source[0]) == source
        events = [r[0] for r in db.execute("SELECT action FROM events WHERE quote_id=?", (q["id"],))]
        assert "REPLY_INSERTED" in events and not any("SENT" in action for action in events)


def test_existing_10_database_is_upgraded_without_losing_data(env):
    c, boss, staff, store, _ = env
    q, body = draft(c, staff, key="legacy-quote-001")
    # A 1.0 database has all existing tables except quote_sources.
    with store.connect() as db:
        db.execute("DROP TABLE quote_sources")
    upgraded = create_app(store.directory)
    with TestClient(upgraded) as after:
        found = after.get("/api/quotes/" + q["id"], headers=staff)
        assert found.status_code == 200 and found.json()["source"] is None
        assert found.json()["total"] == "370.00"
        repeated = after.post("/api/quotes", headers=staff, json=body)
        assert repeated.status_code == 200 and repeated.json()["id"] == q["id"]
        assert after.get("/api/catalog", headers=boss).json()["products"]


def test_whole_cycle_and_archive(env):
    c, boss, staff, store, _ = env
    q, body = draft(c, staff)
    assert q["status"] == "DRAFT" and q["total"] == "370.00"
    assert c.get("/api/quotes/"+q["id"]+"/pdf", headers=staff).status_code == 409
    q = submit(c, staff, q)
    inbox = c.get("/api/notifications", headers=boss).json()
    assert len(inbox) == 1 and inbox[0]["quote_id"] == q["id"]
    assert c.post("/api/quotes/"+q["id"]+"/approve", headers=staff, json={"expected_revision": 1}).status_code == 403
    r = c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1})
    assert r.status_code == 200, r.text
    approved = r.json()
    assert approved["status"] == "COMPLETED", approved
    assert "370.00" in approved["approval"]["replies"]["zh"]
    pdf = c.get("/api/quotes/"+q["id"]+"/pdf", headers=staff)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert not c.get("/api/notifications", headers=boss).json()
    ledger = c.get("/api/exports/ledger", headers=staff)
    book = load_workbook(io.BytesIO(ledger.content), data_only=False)
    assert book.sheetnames == ["Quotes", "Items", "Activity"]
    assert book["Quotes"].max_row == 2
    assert book["Quotes"]["E2"].value == 370
    assert book["Items"]["J2"].value == 370
    assert any(r[5].value == "APPROVED" for r in book["Activity"].iter_rows(min_row=2))
    assert any(r[5].value == "EXPORT_COMPLETED" for r in book["Activity"].iter_rows(min_row=2))


def test_repeat_requests_do_not_duplicate(env):
    c, boss, staff, store, _ = env
    q, body = draft(c, staff)
    assert c.post("/api/quotes", headers=staff, json=body).json()["id"] == q["id"]
    submit(c, staff, q); submit(c, staff, q)
    for _ in range(2):
        assert c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1}).json()["status"] == "COMPLETED"
    with store.connect() as db:
        assert db.execute("SELECT count(*) FROM approvals").fetchone()[0] == 1
        assert db.execute("SELECT count(*) FROM notifications").fetchone()[0] == 1
        assert db.execute("SELECT count(*) FROM events WHERE action='APPROVED'").fetchone()[0] == 1
    body["items"][0]["quantity"] = "30"
    assert c.post("/api/quotes", headers=staff, json=body).status_code == 409


def test_revision_reapproval_and_old_pdf_retained(env):
    c, boss, staff, store, upload = env
    q, body = draft(c, staff); submit(c, staff, q)
    c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1})
    old_pdf = c.get("/api/quotes/"+q["id"]+"/pdf?revision=1", headers=staff).content
    body["items"][0]["quantity"] = "30";body["expected_revision"] = 1
    revised = c.put("/api/quotes/"+q["id"], headers=staff, json=body).json()
    assert revised["revision"] == 2 and revised["status"] == "DRAFT" and revised["approval"] is None
    assert revised["total"] == "555.00"
    assert c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1}).status_code == 409
    assert c.get("/api/quotes/"+q["id"]+"/pdf", headers=staff).status_code == 409
    assert c.get("/api/quotes/"+q["id"]+"/pdf?revision=1", headers=staff).content == old_pdf
    submit(c, staff, revised)
    assert c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 2}).json()["status"] == "COMPLETED"
    with store.connect() as db:
        assert db.execute("SELECT count(*) FROM approvals").fetchone()[0] == 2


def test_export_failure_then_retry(env):
    c, boss, staff, store, _ = env
    q, _ = draft(c, staff); submit(c, staff, q)
    with patch("backend.app.create_ledger", side_effect=PermissionError("test: workbook open")):
        response = c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1})
    assert response.json()["status"] == "EXPORT_PENDING"
    assert response.json()["approval"]["approved_by"] == "boss"
    retry = c.post("/api/quotes/"+q["id"]+"/export", headers=staff, json={"expected_revision": 1})
    assert retry.json()["status"] == "COMPLETED", retry.text
    with store.connect() as db:
        assert db.execute("SELECT count(*) FROM approvals").fetchone()[0] == 1
        assert db.execute("SELECT count(*) FROM events WHERE action='EXPORT_FAILED'").fetchone()[0] == 1


def test_return_and_permission_boundaries(env):
    c, boss, staff, store, _ = env
    assert c.get("/api/catalog").status_code == 401
    assert c.put("/api/shop", headers=staff, json={"name": "X"}).status_code == 403
    q, _ = draft(c, staff); submit(c, staff, q)
    rejected = c.post("/api/quotes/"+q["id"]+"/reject", headers=boss, json={"expected_revision": 1, "reason": "确认数量"}).json()
    assert rejected["status"] == "REJECTED"
    assert c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1}).status_code == 409
    assert c.get("/api/quotes/"+q["id"]+"/pdf", headers=staff).status_code == 409
    assert c.get("/api/health", headers={"Origin": "https://evil.example"}).status_code == 403
    assert c.get("/api/health", headers={"Host": "evil.example"}).status_code == 400


def test_catalog_invalid_rows_and_column_mapping(env):
    c, boss, staff, store, upload = env
    assert c.post("/api/catalog/preview", headers=staff, json=upload).status_code == 403
    book = load_workbook(io.BytesIO(base64.b64decode(upload["content_b64"])))
    sheet = book.active
    sheet["A3"] = sheet["A2"].value
    buffer = io.BytesIO();book.save(buffer)
    bad = {**upload, "content_b64": base64.b64encode(buffer.getvalue()).decode()}
    preview = c.post("/api/catalog/preview", headers=boss, json=bad).json()
    assert not preview["valid"] and any("Duplicate" in e["message"] for e in preview["errors"])
    assert c.post("/api/catalog/import", headers=boss, json=bad).status_code == 422
    sheet["A3"] = "CEM-A-25";sheet["H2"] = "=1+1"
    buffer = io.BytesIO();book.save(buffer);bad["content_b64"] = base64.b64encode(buffer.getvalue()).decode()
    assert not c.post("/api/catalog/preview", headers=boss, json=bad).json()["valid"]
    sheet["H2"] = 18.50;sheet["H1"] = "My Price"
    buffer = io.BytesIO();book.save(buffer)
    mapped = {**upload, "content_b64": base64.b64encode(buffer.getvalue()).decode(), "mapping": {"unit_price": "My Price"}}
    assert c.post("/api/catalog/preview", headers=boss, json=mapped).json()["valid"]


def test_untrusted_prices_rejected_and_formula_text_safe(env):
    c, boss, staff, store, _ = env
    q, body = draft(c, staff, customer="=HYPERLINK(\"https://example.com\")")
    body["request_key"] = "test-key-evil";body["items"][0]["unit_price"] = "0.01"
    assert c.post("/api/quotes", headers=staff, json=body).status_code == 422
    submit(c, staff, q)
    c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1})
    workbook = load_workbook(io.BytesIO(c.get("/api/exports/ledger", headers=staff).content))
    cell = workbook["Quotes"]["C2"]
    assert cell.data_type != "f" and cell.value.startswith("'=")


def test_snapshot_prices_survive_catalog_change(env):
    c, boss, staff, store, upload = env
    q, _ = draft(c, staff);submit(c, staff, q)
    c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1})
    book = load_workbook(io.BytesIO(base64.b64decode(upload["content_b64"])))
    book.active["H2"] = 25.00
    buffer = io.BytesIO();book.save(buffer);upload["content_b64"] = base64.b64encode(buffer.getvalue()).decode()
    assert c.post("/api/catalog/import", headers=boss, json=upload).status_code == 200
    old = c.get("/api/quotes/"+q["id"], headers=staff).json()
    assert old["total"] == "370.00" and old["items"][0]["unit_price"] == "18.50"
    new, _ = draft(c, staff, key="new-price-test")
    assert new["total"] == "500.00"


def test_pdf_pagination_for_maximum_lines(env):
    c, boss, staff, store, _ = env
    q, body = draft(c, staff)
    body["items"] = [{"sku": "CEM-A-50", "quantity": "1"} for _ in range(40)]
    body["request_key"] = "forty-lines-test"
    response = c.post("/api/quotes", headers=staff, json=body)
    assert response.status_code == 200, response.text
    q = response.json()
    assert q["total"] == "740.00"
    submit(c, staff, q)
    approved = c.post("/api/quotes/"+q["id"]+"/approve", headers=boss, json={"expected_revision": 1}).json()
    assert approved["status"] == "COMPLETED", approved
    pdf = c.get("/api/quotes/"+q["id"]+"/pdf", headers=staff)
    import re
    assert len(re.findall(rb"/Type\s*/Page\b", pdf.content)) >= 3


def test_multiple_api_validation_failures_do_not_mutate_data(env):
    c, boss, staff, store, upload = env
    for value in ["0", "-1", "NaN", "Infinity", "1.5", "1000001"]:
        response = c.post("/api/quotes", headers=staff, json={
            "customer": "Demo", "inquiry": "test", "language": "en",
            "items": [{"sku": "CEM-A-50", "quantity": value}],
            "source_confirmed": True, "request_key": "invalid-"+value})
        assert response.status_code == 422, response.text
    bad_sheet = {**upload, "sheet_name": "missing-sheet"}
    assert c.post("/api/catalog/preview", headers=boss, json=bad_sheet).status_code == 422
    with store.connect() as db:
        assert db.execute("SELECT count(*) FROM quotes").fetchone()[0] == 0
