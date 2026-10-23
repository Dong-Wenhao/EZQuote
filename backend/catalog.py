"""Strict XLSX import and a conservative catalog-based trilingual parser."""
import base64
import hashlib
import io
import re
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from openpyxl import load_workbook

FIELDS = ["sku", "name_zh", "name_en", "name_ms", "brand", "spec", "unit", "unit_price", "aliases",
          "category", "applications", "preferences", "unit_cost"]
HEADER_ALIASES = {
    "sku": ["sku", "商品编码", "货号", "kod"],
    "name_zh": ["name_zh", "中文名称", "商品名称", "商品"],
    "name_en": ["name_en", "英文名称", "product", "product_name"],
    "name_ms": ["name_ms", "马来语名称", "nama"],
    "brand": ["brand", "品牌", "jenama"],
    "spec": ["spec", "规格", "specification", "spesifikasi"],
    "unit": ["unit", "单位"],
    "unit_price": ["unit_price", "单价", "price", "harga"],
    "aliases": ["aliases", "别名", "alias"],
    "category": ["category", "类型", "产品类型", "kategori"],
    "applications": ["applications", "用途", "application", "kegunaan"],
    "preferences": ["preferences", "特性", "preference", "ciri"],
    "unit_cost": ["unit_cost", "成本单价", "cost", "kos"],
}
UNITS = {"bag": ["bags", "bag", "beg", "袋", "包"],
         "piece": ["pieces", "piece", "pcs", "pc", "batang", "unit", "个", "支", "根", "件"],
         "roll": ["rolls", "roll", "gulung", "卷"], "kg": ["kg", "公斤", "千克"]}


class CatalogError(ValueError):
    pass


def norm(text):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(text or "")).lower()).strip()


def compact(text):
    return re.sub(r"[\s_\-]+", "", norm(text))


def decimal_value(value, kind="price"):
    try:
        n = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        raise CatalogError(f"Invalid {kind}")
    if not n.is_finite() or n <= 0 or n > Decimal("1000000"):
        raise CatalogError(f"{kind} must be greater than zero and at most 1000000")
    places = max(0, -n.as_tuple().exponent)
    if places > (2 if kind == "price" else 3):
        raise CatalogError(f"Too many decimal places in {kind}")
    return n


def amount(value):
    return str(Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def read_catalog(content_b64, mapping=None, sheet_name=None):
    try:
        raw = base64.b64decode(content_b64, validate=True)
    except Exception:
        raise CatalogError("Invalid Excel upload")
    if len(raw) > 5 * 1024 * 1024:
        raise CatalogError("Excel file exceeds 5 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist()) > 500 or sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                raise CatalogError("Excel file is too large after decompression")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except CatalogError:
        raise
    except Exception:
        raise CatalogError("Please upload a valid .xlsx workbook")
    try:
        if sheet_name and sheet_name not in book.sheetnames:
            raise CatalogError("Worksheet not found")
        sheet = book[sheet_name] if sheet_name else book.worksheets[0]
        if (sheet.max_row or 0) > 2001 or (sheet.max_column or 0) > 32:
            raise CatalogError("Use at most 2000 products and 32 columns")
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index > 2000 or len(row) > 32:
                raise CatalogError("Use at most 2000 products and 32 columns")
            rows.append(row)
        if not rows:
            raise CatalogError("Excel sheet is empty")
        headers = [str(x or "").strip() for x in rows[0]]
        if len(set(headers)) != len(headers) or any(not h for h in headers):
            raise CatalogError("Every column in the first row needs a unique header")
        inferred = {field: next((h for h in headers if norm(h) in aliases), "") for field, aliases in HEADER_ALIASES.items()}
        mapped = {**inferred, **(mapping or {})}
        indexes = {k: headers.index(v) for k, v in mapped.items() if k in FIELDS and v in headers}
        errors = []
        required = [k for k in ["sku", "brand", "spec", "unit", "unit_price"] if k not in indexes]
        if not any(k in indexes for k in ["name_zh", "name_en", "name_ms"]):
            required.append("name_zh / name_en / name_ms")
        if required:
            errors.append({"row": 1, "message": "Missing columns: " + ", ".join(required)})
            return {"headers": headers, "mapping": mapped, "sheets": book.sheetnames,
                    "sheet": sheet.title, "errors": errors, "products": [], "valid": False}, raw
        products = []
        seen = set()
        for row_num, row in enumerate(rows[1:], 2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            def val(key):
                return row[indexes[key]] if key in indexes and indexes[key] < len(row) else ""
            try:
                p = {k: str(val(k) or "").strip() for k in FIELDS if k not in {"unit_price", "unit_cost"}}
                if any(len(x) > 500 for x in p.values()):
                    raise CatalogError("Cell text exceeds 500 characters")
                if not p["sku"] or not p["brand"] or not p["spec"]:
                    raise CatalogError("SKU, brand and specification are required")
                if p["sku"].lower() in seen:
                    raise CatalogError("Duplicate SKU")
                if not any(p[k] for k in ["name_zh", "name_en", "name_ms"]):
                    raise CatalogError("A product name is required")
                unit = norm(p["unit"])
                p["unit"] = next((k for k, aliases in UNITS.items() if unit == k or unit in aliases), "")
                if not p["unit"]:
                    raise CatalogError("Unit must be bag / piece / roll / kg")
                p["unit_price"] = amount(decimal_value(val("unit_price")))
                if str(val("unit_cost") or "").strip():
                    p["unit_cost"] = amount(decimal_value(val("unit_cost")))
                seen.add(p["sku"].lower())
                products.append(p)
            except CatalogError as e:
                errors.append({"row": row_num, "message": str(e)})
        if not products and not errors:
            errors.append({"row": 2, "message": "No products found"})
        return {"headers": headers, "mapping": mapped, "sheets": book.sheetnames,
                "sheet": sheet.title, "products": products, "errors": errors,
                "valid": not errors, "sha256": hashlib.sha256(raw).hexdigest()}, raw
    finally:
        book.close()


def detect_language(text):
    if re.search(r"[\u3400-\u9fff]", text):
        return "zh"
    return "ms" if re.search(r"\b(simen|berapa|harga|beg|nak|mahu|batang|tolong|gulung|jenama)\b", norm(text)) else "en"


def contains(text, alias):
    alias = norm(alias)
    if not alias:
        return False
    if re.search(r"[\u3400-\u9fff]", alias):
        return compact(alias) in compact(text)
    return bool(re.search(r"(?<![a-z0-9])" + re.escape(alias).replace(r"\ ", r"\s*") + r"(?![a-z0-9])", norm(text)))


def parse_inquiry(text, products):
    # Import here because the optional AI extractor also imports this catalog module.
    from .requirement_tokens import quantities as scan_quantities
    language = detect_language(text)
    lines = [s.strip() for s in re.split(r"[\n;；]+", text) if s.strip()]
    if not lines:
        raise CatalogError("Inquiry is empty")
    output = []
    for line in lines:
        # Ignore greetings, but never discard a line containing numbers or product terms.
        if len(lines) > 1 and re.fullmatch(r"(?:hi|hello|你好|您好|hai|salam)[ !！,.，。]*", norm(line)):
            continue
        sku_matches = [p for p in products if contains(line, p["sku"])]
        candidates = sku_matches or [p for p in products if any(contains(line, a)
                      for a in [p["name_zh"], p["name_en"], p["name_ms"], *re.split(r"[|；;]", p["aliases"])])]
        reasons = []
        if not sku_matches:
            brands = sorted({p["brand"] for p in candidates if contains(line, p["brand"])})
            # Brand aliases such as A牌 and Brand A are represented in the aliases column.
            brand_tokens = re.findall(r"(?:brand|jenama)\s*([a-z0-9]+)|([a-z0-9]+)\s*牌", norm(line))
            if brand_tokens:
                tokens = {compact(a or b) for a, b in brand_tokens}
                candidates = [p for p in candidates if compact(p["brand"]).replace("brand", "") in tokens]
                brands = sorted({p["brand"] for p in candidates})
            elif brands:
                candidates = [p for p in candidates if p["brand"] in brands]
            else:
                reasons.append("brand")
            specs = re.findall(r"\d+(?:\.\d+)?\s*(?:kg|公斤|千克|mm|cm|ml|litre|liter|l|m)(?![a-z])", norm(line))
            if specs:
                candidates = [p for p in candidates if compact(p["spec"]) in [compact(s) for s in specs]]
            else:
                # Require a specification unless the customer used an exact SKU.
                reasons.append("spec")
        mentions = scan_quantities(line, UNITS, allow_kg=bool(candidates and all(p["unit"] == "kg" for p in candidates)))
        quantities = [(m['value'] or '', m['unit']) for m in mentions]
        if len(quantities) != 1:
            reasons.append("quantity" if not quantities else "multiple_quantities")
        quantity, unit = quantities[0] if len(quantities) == 1 else ("", "")
        for mention in mentions:
            if mention['error']:
                reasons.append('multiple_quantities' if mention['error']=='multiple_quantities' else 'invalid_quantity')
        if quantity:
            try:
                q = decimal_value(quantity, "quantity")
                if unit != "kg" and q != q.to_integral_value():
                    raise CatalogError("Whole units are required")
            except CatalogError:
                reasons.append("invalid_quantity")
        if candidates and unit:
            candidates = [p for p in candidates if p["unit"] == unit]
        if not candidates:
            reasons.append("unknown_product")
        elif len(candidates) > 1:
            reasons.append("ambiguous_product")
        reasons = list(dict.fromkeys(reasons))
        selected = candidates[0] if len(candidates) == 1 and not reasons else None
        questions = make_questions(reasons, language)
        output.append({"source_text": line, "sku": selected["sku"] if selected else "",
                       "quantity": quantity, "unit": unit, "candidates": candidates, "quantity_mentions": mentions,
                       "issues": reasons, "questions": questions})
    ready = bool(output) and all(not i["issues"] for i in output)
    return {"language": language, "status": "READY" if ready else "NEEDS_INFO", "items": output,
            "clarification": "\n".join(q for i in output for q in i["questions"]),
            "method": "catalog_rules"}


def make_questions(reasons, language):
    wording = {
        "zh": {"brand": "请问需要哪个品牌？", "spec": "请补充包装或规格，例如 50kg。", "quantity": "请补充数量和单位，例如 20 袋。", "multiple_quantities": "请把每种商品分别写一行，并确认各自数量。", "invalid_quantity": "请提供有效的正数数量，整件商品使用整数。", "unknown_product": "商品或规格未匹配，请确认名称、品牌和规格。", "ambiguous_product": "有多个候选商品，请确认具体规格或 SKU。"},
        "en": {"brand": "Which brand do you need?", "spec": "Please confirm the specification, e.g. 50kg.", "quantity": "Please confirm quantity and unit, e.g. 20 bags.", "multiple_quantities": "Please put each product on a separate line and confirm its quantity.", "invalid_quantity": "Please provide a positive quantity; use whole numbers for whole units.", "unknown_product": "Please confirm the product name, brand and specification.", "ambiguous_product": "Several products match. Please confirm the exact specification or SKU."},
        "ms": {"brand": "Jenama apa yang diperlukan?", "spec": "Sila sahkan spesifikasi, contohnya 50kg.", "quantity": "Sila sahkan kuantiti dan unit, contohnya 20 beg.", "multiple_quantities": "Sila tulis setiap produk pada baris berasingan dan sahkan kuantitinya.", "invalid_quantity": "Sila berikan kuantiti positif; gunakan nombor bulat untuk unit penuh.", "unknown_product": "Sila sahkan nama produk, jenama dan spesifikasi.", "ambiguous_product": "Beberapa produk sepadan. Sila sahkan spesifikasi atau SKU yang tepat."}}
    return [wording[language][r] for r in reasons]


def price_items(entries, products):
    if not 1 <= len(entries) <= 40:
        raise CatalogError("Use between 1 and 40 quotation lines")
    by_sku = {p["sku"]: p for p in products}
    result = []
    for item in entries:
        p = by_sku.get(item["sku"])
        if not p:
            raise CatalogError("Unknown SKU: " + item["sku"])
        qty = decimal_value(item["quantity"], "quantity")
        if p["unit"] != "kg" and qty != qty.to_integral_value():
            raise CatalogError("Whole units are required for " + p["sku"])
        result.append({**p, "quantity": str(qty), "line_total": amount(qty * Decimal(p["unit_price"]))})
    return result, amount(sum(Decimal(i["line_total"]) for i in result))
