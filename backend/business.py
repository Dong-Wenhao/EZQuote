"""Company-owned product, inventory, delivery and outcome data. No external services."""
import base64
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .catalog import CatalogError, amount, compact, read_catalog
from .store import dumps, now

DEFAULT_RULES = {"min_margin_pct": "15.00", "max_discount_pct": "10.00",
                 "max_delivery_subsidy": "3.00", "minimum_similar_orders": 3,
                 "sales_window_days": 90}


def number(value, label, maximum=1_000_000, whole=False):
    if isinstance(value, bool):
        raise CatalogError(f"{label}: a number is required")
    try:
        n = Decimal(str(value).strip())
    except (InvalidOperation, TypeError, ValueError):
        raise CatalogError(f"{label}: a number is required") from None
    if not n.is_finite() or n < 0 or n > maximum:
        raise CatalogError(f"{label}: use a number from 0 to {maximum}")
    if -n.as_tuple().exponent > (0 if whole else 3):
        raise CatalogError(f"{label}: too many decimal places")
    return n


def text(value, label, limit=500):
    result = str(value if value is not None else "").strip()
    if len(result) > limit or result.startswith("="):
        raise CatalogError(f"{label}: invalid or excessively long text")
    return result


def read_business(content_b64):
    """Validate all sheets before the single database transaction imports anything."""
    products, raw = read_catalog(content_b64, sheet_name="Products")
    if not products["valid"]:
        details = "; ".join(f"row {x['row']}: {x['message']}" for x in products["errors"][:8])
        raise CatalogError("Products: " + details)
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        for sheet in ("Inventory", "History", "Delivery", "Rules"):
            if sheet not in book.sheetnames:
                raise CatalogError(f"Missing worksheet: {sheet}")

        def rows(name, fields, maximum=10000):
            sheet = book[name]
            if (sheet.max_row or 0) > maximum + 1 or (sheet.max_column or 0) > 24:
                raise CatalogError(f"{name}: too many rows or columns")
            it = sheet.iter_rows(values_only=True)
            headers = [str(v or "").strip() for v in next(it, [])]
            if not headers or len(set(headers)) != len(headers) or any(not h for h in headers):
                raise CatalogError(f"{name}: unique nonempty headers are required")
            if any(f not in headers for f in fields):
                raise CatalogError(f"{name}: required columns: {', '.join(fields)}")
            result = []
            for line, row in enumerate(it, 2):
                if line > maximum + 1:
                    raise CatalogError(f"{name}: too many rows")
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                if any(isinstance(v, str) and v.startswith("=") for v in row):
                    raise CatalogError(f"{name} row {line}: formulas are not accepted as source data")
                result.append(({h: row[j] if j < len(row) else None for j, h in enumerate(headers)}, line))
            return result

        by_sku = {p["sku"]: p for p in products["products"]}
        inventory = {}
        for row, line in rows("Inventory", ("sku", "stock", "low_stock", "high_stock"), 2000):
            sku = text(row["sku"], "Inventory sku", 100)
            if sku not in by_sku or sku in inventory:
                raise CatalogError(f"Inventory row {line}: unknown or duplicate SKU {sku}")
            quantities = {key: str(number(row[key], "Inventory " + key)) for key in ("stock", "low_stock", "high_stock")}
            if by_sku[sku]["unit"] != "kg" and any(Decimal(v) != Decimal(v).to_integral_value() for v in quantities.values()):
                raise CatalogError(f"Inventory row {line}: whole units are required")
            if Decimal(quantities["high_stock"]) <= Decimal(quantities["low_stock"]):
                raise CatalogError(f"Inventory row {line}: high_stock must exceed low_stock")
            inventory[sku] = quantities
        missing = set(by_sku) - set(inventory)
        if missing:
            raise CatalogError("Inventory: missing SKUs: " + ", ".join(sorted(missing)[:10]))
        for p in by_sku.values():
            if not p.get("category") or not p.get("unit_cost"):
                raise CatalogError(f"Products {p['sku']}: category and positive unit_cost are required for pricing")

        delivery = []
        occupied = set()
        for row, line in rows("Delivery", ("area", "aliases", "base_fee", "cost"), 2000):
            area = text(row["area"], "Delivery area", 160)
            aliases = text(row["aliases"], "Delivery aliases")
            keys = {compact(a) for a in [area, *aliases.split("|")] if a.strip()}
            if not area or keys & occupied:
                raise CatalogError(f"Delivery row {line}: empty or overlapping area aliases")
            occupied |= keys
            delivery.append({"area": area, "aliases": aliases,
                             "base_fee": amount(number(row["base_fee"], "base_fee")),
                             "cost": amount(number(row["cost"], "delivery cost"))})

        rules = dict(DEFAULT_RULES)
        rule_seen = set()
        for row, line in rows("Rules", ("key", "value"), 100):
            key = text(row["key"], "Rule key", 100)
            if key not in DEFAULT_RULES or key in rule_seen:
                raise CatalogError(f"Rules row {line}: unknown or duplicate rule {key}")
            rule_seen.add(key)
            if key in {"minimum_similar_orders", "sales_window_days"}:
                rules[key] = int(number(row["value"], key, maximum=3650, whole=True))
                if rules[key] < 1:
                    raise CatalogError(f"{key} must be positive")
            else:
                rules[key] = amount(number(row["value"], key, maximum=100 if key.endswith("pct") else 1000))
        if Decimal(rules["min_margin_pct"]) >= 100 or Decimal(rules["max_discount_pct"]) >= 100:
            raise CatalogError("Margin and discount percentages must be less than 100")

        history = []
        history_keys, order_values = set(), {}
        for row, line in rows("History", ("order_id", "date", "sku", "quantity", "area", "discount_pct", "delivery_cost", "delivery_fee", "outcome")):
            order_id, sku = text(row["order_id"], "order_id", 100), text(row["sku"], "history sku", 100)
            if not order_id or sku not in by_sku or (order_id, sku) in history_keys:
                raise CatalogError(f"History row {line}: invalid or duplicate order/SKU")
            history_keys.add((order_id, sku))
            raw_date = row["date"]
            try:
                day = raw_date.date() if isinstance(raw_date, datetime) else raw_date if isinstance(raw_date, date) else date.fromisoformat(str(raw_date)[:10])
            except (ValueError, TypeError):
                raise CatalogError(f"History row {line}: use an ISO date YYYY-MM-DD") from None
            if day > date.today():
                raise CatalogError(f"History row {line}: future outcomes are not historical evidence")
            quantity = number(row["quantity"], "history quantity")
            if quantity <= 0 or (by_sku[sku]["unit"] != "kg" and quantity != quantity.to_integral_value()):
                raise CatalogError(f"History row {line}: invalid quantity")
            outcome = text(row["outcome"], "outcome").lower()
            if outcome not in {"won", "lost", "open"}:
                raise CatalogError(f"History row {line}: outcome must be won / lost / open")
            area = text(row["area"], "history area", 160)
            area = next((d["area"] for d in delivery if compact(area) in {compact(a) for a in [d["area"], *d["aliases"].split("|")] if a}), area)
            discount = amount(number(row["discount_pct"], "discount_pct", maximum=100))
            shipping_cost = amount(number(row["delivery_cost"], "delivery_cost"))
            shipping_fee = amount(number(row["delivery_fee"], "delivery_fee"))
            common = (day.isoformat(), compact(area), discount, shipping_cost, shipping_fee, outcome)
            if order_id in order_values and order_values[order_id] != common:
                raise CatalogError(f"History row {line}: rows for one order must share date, area, discount, freight and outcome")
            order_values[order_id] = common
            history.append({"order_id": order_id, "date": day.isoformat(), "sku": sku, "quantity": str(quantity),
                            "area": area, "discount_pct": discount, "delivery_cost": shipping_cost,
                            "delivery_fee": shipping_fee, "outcome": outcome, "source": "import"})
        return {"products": products["products"], "inventory": inventory, "history": history,
                "delivery": delivery, "rules": rules,
                "fingerprint": hashlib.sha256(raw).hexdigest()}, raw
    finally:
        book.close()


def get_business(store, db=None):
    if db is None:
        with store.connect() as conn:
            return get_business(store, conn)
    row = db.execute("SELECT * FROM business_versions ORDER BY id DESC LIMIT 1").fetchone()
    if not row:
        return {"id": 0, "configured": False, "inventory": {}, "history": [], "delivery": [], "rules": dict(DEFAULT_RULES)}
    return {**json.loads(row["data"]), "id": row["id"], "catalog_id": row["catalog_id"],
            "configured": True, "created_at": row["created_at"], "filename": row["filename"]}


def import_business(store, data, raw, filename, actor):
    original = store.directory / "uploads" / (data["fingerprint"] + ".xlsx")
    if not original.exists():
        original.write_bytes(raw)
    with store.transaction() as db:
        current = db.execute("SELECT * FROM business_versions ORDER BY id DESC LIMIT 1").fetchone()
        if current and current["fingerprint"] == data["fingerprint"] and store.catalog(db)["id"] == current["catalog_id"]:
            return get_business(store, db)
        safe = Path(filename.replace("\\", "/")).name
        cat = db.execute("INSERT INTO catalogs(created_at,filename,products) VALUES(?,?,?)", (now(), safe, dumps(data["products"]))).lastrowid
        db.execute("INSERT INTO business_versions(created_at,filename,fingerprint,catalog_id,data) VALUES(?,?,?,?,?)",
                   (now(), safe, data["fingerprint"], cat, dumps({k: v for k, v in data.items() if k not in {"products", "fingerprint"}})))
        store.event(db, {"id": "BUSINESS", "revision": 1}, actor, "BUSINESS_IMPORTED",
                    {"products": len(data["products"]), "orders": len({h['order_id'] for h in data['history']})})
        result = get_business(store, db)
    return result


def historical_rows(store, db=None, exclude_quote=None):
    """One latest evaluated approval per quote; never count revisions as extra orders."""
    if db is None:
        with store.connect() as conn:
            return historical_rows(store, conn, exclude_quote)
    imported = get_business(store, db)["history"]
    live = []
    seen = set()
    rows = db.execute("SELECT o.*,a.snapshot FROM quote_outcomes o JOIN approvals a ON a.quote_id=o.quote_id AND a.revision=o.revision ORDER BY o.revision DESC").fetchall()
    for row in rows:
        qid = row["quote_id"]
        if qid in seen or qid == exclude_quote:
            continue
        snap = json.loads(row["snapshot"])
        pricing = snap.get("pricing")
        if not pricing or not pricing.get("decision"):
            continue
        seen.add(qid)
        decision = pricing["decision"]
        for item in snap["items"]:
            live.append({"order_id": qid, "date": row["updated_at"][:10], "sku": item["sku"],
                         "quantity": item["quantity"], "area": pricing["requirement"].get("area", ""),
                         "discount_pct": decision["discount_pct"],
                         "delivery_cost": row["actual_delivery_cost"] or "0.00",
                         "delivery_cost_known": bool(row["actual_delivery_cost"]),
                         "delivery_fee": decision["freight"], "outcome": row["outcome"], "source": "application"})
    return [h for h in imported if h["order_id"] not in seen and h["order_id"] != exclude_quote] + live


def business_summary(store, db=None):
    if db is None:
        with store.connect() as conn:
            return business_summary(store, conn)
    data = get_business(store, db)
    history = historical_rows(store, db)
    orders = {h["order_id"]: h for h in history}
    closed = [h for h in orders.values() if h["outcome"] in {"won", "lost"}]
    won = sum(h["outcome"] == "won" for h in closed)
    return {"configured": data["configured"], "id": data["id"], "catalog_id": data.get("catalog_id", 0),
            "created_at": data.get("created_at"), "filename": data.get("filename", ""),
            "inventory": data["inventory"], "delivery": data["delivery"], "rules": data["rules"],
            "orders": len(orders), "closed_orders": len(closed), "won_orders": won,
            "lost_orders": len(closed) - won, "live_orders": sum(h["source"] == "application" for h in orders.values()),
            "conversion_rate": str((Decimal(won) / len(closed)).quantize(Decimal("0.0001"))) if closed else None}
