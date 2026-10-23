"""Build every outward-facing result from a boss-approved snapshot."""
import io
import json
import os
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]


def pdf_font():
    if "QuoteFont" in pdfmetrics.getRegisteredFontNames():
        return "QuoteFont"
    path = ROOT / "assets" / "NotoSansSC.ttf"
    if path.exists():
        pdfmetrics.registerFont(TTFont("QuoteFont", str(path)))
        return "QuoteFont"
    # Supported by ReportLab, but package QA requires the embedded font above.
    if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def product_name(item, language):
    return item.get("name_" + language) or item.get("name_en") or item.get("name_zh") or item["sku"]


def reply_text(snapshot, language):
    units = {"zh": {"bag": "袋", "piece": "件", "roll": "卷", "kg": "公斤"},
             "en": {"bag": "bags", "piece": "pieces", "roll": "rolls", "kg": "kg"},
             "ms": {"bag": "beg", "piece": "unit", "roll": "gulung", "kg": "kg"}}
    rows = [f"{i['brand']} {product_name(i, language)} {i['spec']} × {i['quantity']} {units[language][i['unit']]}" for i in snapshot["items"]]
    detail = ("；" if language == "zh" else "; ").join(rows)
    total = snapshot["total"]
    quote_id = snapshot["id"]
    pricing = snapshot.get("pricing")
    if pricing and pricing.get("decision"):
        decision, requirement = pricing["decision"], pricing["requirement"]
        area = requirement["area"] if requirement["fulfilment"] == "delivery" else {"zh": "自提", "en": "Collection", "ms": "Ambil sendiri"}[language]
        parts = {
            "zh": f"原商品金额 RM{pricing['recommendation']['original_product_price']}，折扣 {decision['discount_pct']}%（减 RM{decision['discount_amount']}），运费 RM{decision['freight']}。配送方式/地区：{area}。",
            "en": f"Products RM{pricing['recommendation']['original_product_price']}, discount {decision['discount_pct']}% (RM{decision['discount_amount']} off), delivery RM{decision['freight']}. Delivery/collection: {area}.",
            "ms": f"Produk RM{pricing['recommendation']['original_product_price']}, diskaun {decision['discount_pct']}% (tolak RM{decision['discount_amount']}), penghantaran RM{decision['freight']}. Penghantaran/ambil: {area}."}
        if language == "zh":
            return f"您好，{detail}。{parts[language]}最终合计 RM{total}。报价编号 {quote_id}，详情请见报价单。"
        if language == "ms":
            return f"Salam, {detail}. {parts[language]} Jumlah akhir RM{total}. Rujukan {quote_id}. Sila lihat sebut harga yang dilampirkan."
        return f"Hello, {detail}. {parts[language]} Final total RM{total}. Reference {quote_id}. Please see the attached quotation."
    if language == "zh":
        return f"您好，{detail}，报价合计 RM{total}。报价编号 {quote_id}，详情请见报价单。"
    if language == "ms":
        return f"Salam, {detail}. Jumlah sebut harga ialah RM{total}. Rujukan {quote_id}. Sila lihat sebut harga yang dilampirkan."
    return f"Hello, {detail}. The quotation total is RM{total}. Reference {quote_id}. Please see the attached quotation."


def create_pdf(snapshot, output_path):
    output_path = Path(output_path)
    temporary = output_path.with_suffix(".tmp.pdf")
    font = pdf_font()
    normal = ParagraphStyle("body", fontName=font, fontSize=9, leading=14, wordWrap="CJK", textColor=colors.HexColor("#152d48"))
    title = ParagraphStyle("title", parent=normal, fontSize=22, leading=30)
    small = ParagraphStyle("small", parent=normal, fontSize=8, leading=12)
    right = ParagraphStyle("right", parent=normal, alignment=TA_RIGHT)
    def p(text, style=normal):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)
    shop = snapshot["shop"]
    body = [p(shop["name"], title), p("QUOTATION / 报价单 / SEBUT HARGA"), Spacer(1, 5 * mm)]
    if shop.get("address"):
        body.append(p(shop["address"]))
    if shop.get("phone"):
        body.append(p(shop["phone"]))
    body += [Spacer(1, 5 * mm), p(f"{snapshot['id']}   ·   Revision {snapshot['revision']}"),
             p("Customer / 客户: " + snapshot["customer"]),
             p("Approved / 批准: " + snapshot["approved_at"][:19].replace("T", " ") + " UTC"), Spacer(1, 5 * mm)]
    white = ParagraphStyle("white", parent=small, textColor=colors.white)
    data = [[p(t, white) for t in ["Product / 商品", "Qty", "Unit", "Price RM", "Amount RM"]]]
    for item in snapshot["items"]:
        description = f"{item['sku']}\n{item['brand']} {product_name(item, snapshot['language'])}\n{item['spec']}"
        data.append([p(description), p(item["quantity"], right), p(item["unit"]),
                     p(item["unit_price"], right), p(item["line_total"], right)])
    table = Table(data, colWidths=[84*mm, 17*mm, 18*mm, 26*mm, 29*mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#174a78")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6fa")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LINEBELOW", (0, 0), (-1, 0), .5, colors.HexColor("#ced9e3")),
    ]))
    body += [table, Spacer(1, 5*mm)]
    pricing = snapshot.get("pricing")
    if pricing and pricing.get("decision"):
        rec, decision, requirement = pricing["recommendation"], pricing["decision"], pricing["requirement"]
        body += [p("Delivery / 配送地区: " + (requirement["area"] if requirement["fulfilment"] == "delivery" else "Collection / 自提")),
                 p("Original products / 商品原价   RM " + rec["original_product_price"], right),
                 p("Base delivery / 基础运费   RM " + rec["base_freight"], right),
                 p("Original total / 原始总额   RM " + rec["original_total"], right),
                 p("Discount / 折扣 " + decision["discount_pct"] + "%   - RM " + decision["discount_amount"], right),
                 p("Final delivery / 最终运费   RM " + decision["freight"], right),
                 p("Approved by / 审批人: " + snapshot["approved_by"])]
    body += [p("FINAL TOTAL / 最终合计   RM " + snapshot["total"], right), Spacer(1, 7*mm)]
    if snapshot["notes"]:
        body += [p("Notes / 备注"), p(snapshot["notes"]), Spacer(1, 4*mm)]
    body += [p(shop.get("terms", ""), small), p("Quotation only. This is not an invoice or a payment receipt.", small)]
    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#60758a"))
        canvas.drawString(18*mm, 12*mm, snapshot["id"] + f" / r{snapshot['revision']}")
        canvas.drawRightString(192*mm, 12*mm, str(doc.page))
    SimpleDocTemplate(str(temporary), pagesize=(210*mm, 297*mm), rightMargin=18*mm,
                      leftMargin=18*mm, topMargin=18*mm, bottomMargin=20*mm).build(body, onFirstPage=footer, onLaterPages=footer)
    os.replace(temporary, output_path)


def literal(value):
    """Never allow user-controlled text to become an Excel formula."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def style_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="174A78")
        cell.alignment = Alignment(vertical="center", horizontal="center")
    sheet.row_dimensions[1].height = 28
    for r in sheet.iter_rows(min_row=2):
        for c in r:
            c.font = Font(name="Arial", size=10, color="152D48")
            c.alignment = Alignment(vertical="center", horizontal="right" if isinstance(c.value, (int, float)) else "left", indent=1)
            if c.row % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F6FA")
    for i, w in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = w


def create_ledger(store, output_path):
    """Rebuild from all immutable approvals, so retries cannot append duplicates."""
    with store.connect() as db:
        approvals = db.execute("SELECT snapshot FROM approvals ORDER BY quote_id, revision").fetchall()
        events = db.execute("SELECT * FROM events ORDER BY id").fetchall()
        outcomes = db.execute("SELECT * FROM quote_outcomes ORDER BY quote_id,revision").fetchall()
    book = Workbook()
    quotes = book.active
    quotes.title = "Quotes"
    quotes.append(["Quote ID", "Revision", "Customer", "Currency", "Total", "Approved by", "Approved at UTC", "Catalog version"])
    lines = book.create_sheet("Items")
    lines.append(["Quote ID", "Revision", "SKU", "Brand", "Product", "Spec", "Unit", "Quantity", "Unit price MYR", "Line total MYR"])
    for row in approvals:
        s = json.loads(row[0])
        quotes.append([s["id"], s["revision"], literal(s["customer"]), "MYR", float(s["total"]), s["approved_by"], s["approved_at"], s["catalog_id"]])
        for item in s["items"]:
            lines.append([s["id"], s["revision"], literal(item["sku"]), literal(item["brand"]), literal(product_name(item, s["language"])),
                          literal(item["spec"]), item["unit"], float(item["quantity"]), float(item["unit_price"]), float(item["line_total"])])
    log = book.create_sheet("Activity")
    log.append(["Event ID", "Quote ID", "Revision", "Time UTC", "Actor", "Action", "Details"])
    for row in events:
        log.append([row["id"], row["quote_id"], row["revision"], row["at"], literal(row["actor"]), row["action"], literal(row["details"])])
    style_sheet(quotes, [30, 10, 28, 12, 16, 18, 28, 18])
    style_sheet(lines, [30, 10, 22, 18, 30, 16, 12, 14, 20, 20])
    style_sheet(log, [12, 30, 10, 28, 18, 24, 60])
    smart = [json.loads(row[0]) for row in approvals if json.loads(row[0]).get("pricing")]
    if smart:
        pricing_sheet = book.create_sheet("Pricing")
        pricing_sheet.append(["Quote ID", "Revision", "Area", "Fulfilment", "Original products", "Base freight", "Original total", "Suggested discount %", "Suggested freight", "Recommended total", "Final discount %", "Final discount amount", "Final freight", "Final total", "Owner decision", "Decision note", "Similar closed orders", "Conversion rate", "Business version", "Recommendation reasons"])
        requirements = book.create_sheet("Requirements")
        requirements.append(["Quote ID", "Revision", "Requirement JSON", "Selected items", "Original inquiry"])
        for snap in smart:
            pr = snap["pricing"]; rec, dec, req = pr["recommendation"], pr["decision"], pr["requirement"]
            pricing_sheet.append([snap["id"], snap["revision"], literal(req["area"]), req["fulfilment"], float(rec["original_product_price"]), float(rec["base_freight"]), float(rec["original_total"]), float(rec["suggested_discount_pct"]), float(rec["suggested_freight"]), float(rec["recommended_total"]), float(dec["discount_pct"]), float(dec["discount_amount"]), float(dec["freight"]), float(dec["final_total"]), dec["mode"], literal(dec["note"]), rec["similar_orders"], float(rec["conversion_rate"]) if rec["conversion_rate"] is not None else None, rec["business_version"], literal("; ".join(rec["reasons"][snap["language"]]))])
            requirements.append([snap["id"], snap["revision"], literal(json.dumps(req, ensure_ascii=False)), literal(json.dumps(req["selected_items"], ensure_ascii=False)), literal(snap["inquiry"])])
        style_sheet(pricing_sheet, [30, 10, 22, 16] + [19]*10 + [18, 32, 22, 20, 18, 65])
        style_sheet(requirements, [30, 10, 70, 60, 65])
        for row in pricing_sheet.iter_rows(min_row=2):
            for index in [4, 5, 6, 8, 9, 11, 12, 13]:
                row[index].number_format = '"RM "#,##0.00'
            row[17].number_format = "0.0%"
    if outcomes:
        results = book.create_sheet("Outcomes")
        results.append(["Quote ID", "Revision", "Customer outcome", "Note", "Recorded by", "Recorded at UTC", "Outcome version", "Actual delivery cost MYR"])
        for r in outcomes:
            results.append([r["quote_id"], r["revision"], r["outcome"], literal(r["note"]), r["actor"], r["updated_at"], r["version"], float(r["actual_delivery_cost"]) if r["actual_delivery_cost"] else None])
        style_sheet(results, [30, 10, 20, 45, 20, 28, 20, 28])
    for row in quotes.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = '"RM "#,##0.00'
    for row in lines.iter_rows(min_row=2, min_col=8, max_col=10):
        row[0].number_format = '#,##0.000' if lines.cell(row[0].row, 7).value == "kg" else '#,##0'
        row[1].number_format = row[2].number_format = '"RM "#,##0.00'
    temporary = Path(output_path).with_suffix(".tmp.xlsx")
    book.save(temporary)
    os.replace(temporary, output_path)
